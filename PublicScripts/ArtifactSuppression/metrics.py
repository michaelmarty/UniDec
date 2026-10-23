"""Panel-level deconvolution metrics and Excel export helpers."""

from pathlib import Path

import numpy as np
from openpyxl import Workbook
from openpyxl.formatting.rule import ColorScaleRule
from openpyxl.styles import Font
from openpyxl.utils import get_column_letter


SHOW_UNISCORE_LABELS = False


CONFIG_FIELDS = (
    "masslb",
    "massub",
    "massbins",
    "startz",
    "endz",
    "mzsig",
    "psig",
    "subbuff",
    "beta",
    "suppression_satellite",
    "suppression_harmonic",
    "suppression_topn",
    "suppression_topx",
    "peakwindow",
    "peakthresh",
)

COLOR_SCALED_METRICS = {
    "apex_error_Da",
    "apex_abs_error_Da",
    "apex_error_ppm",
    "centroid_error_Da",
    "centroid_abs_error_Da",
    "centroid_error_ppm",
    "DScore",
    "main_peak_area",
    "spectrum_area",
    "main_peak_area_fraction",
    "artifact_area_fraction",
    "main_peak_height_fraction",
    "FWHM_Da",
    "mass_resolving_power",
    "detected_peak_count",
    "UniScore",
    "fit_R_squared",
}

METRIC_DEFINITIONS = (
    ("main_peak", "Highest-intensity picked peak in the deconvolved mass spectrum."),
    ("apex_mass_Da", "Mass at the apex of the main picked peak."),
    ("apex_error_Da", "Apex mass minus the reference mass."),
    ("apex_abs_error_Da", "Absolute apex mass error."),
    ("apex_error_ppm", "Signed apex mass error divided by reference mass, in ppm."),
    ("centroid_mass_Da", "Intensity-weighted mass within UniDec's FWHM interval for the main peak."),
    ("centroid_error_Da", "Mass centroid minus the reference mass."),
    ("centroid_abs_error_Da", "Absolute mass centroid error."),
    ("centroid_error_ppm", "Signed mass centroid error divided by reference mass, in ppm."),
    ("DScore", "UniDec combined DScore for the main peak, on a 0-1 scale."),
    ("main_peak_area", "Trapezoidal area within the reported main-peak area window."),
    ("main_peak_area_window_low_Da", "Lower bound of the main-peak area window."),
    ("main_peak_area_window_high_Da", "Upper bound of the main-peak area window."),
    ("spectrum_area", "Trapezoidal area across the complete deconvolved mass spectrum."),
    ("main_peak_area_fraction", "Area within one FWHM on either side of the main apex, divided by the full-spectrum area."),
    ("artifact_area_fraction", "One minus the main-peak area fraction."),
    ("main_peak_height_fraction", "Main picked-peak height divided by the sum of all picked-peak heights."),
    ("FWHM_Da", "Full width at half maximum reported by UniDec for the main peak."),
    ("mass_resolving_power", "Main-peak apex mass divided by its FWHM."),
    ("detected_peak_count", "Number of peaks passing the script's UniDec peak-picking settings."),
    ("UniScore", "Intensity-weighted spectrum-level UniDec score, on a 0-1 scale."),
    ("fit_R_squared", "UniDec's coefficient of determination for the reconvolved fit to the input data."),
)


def _finite_number(value):
    """Return a plain finite Python number, or None for an unavailable value."""
    try:
        number = float(value)
    except (TypeError, ValueError):
        return None
    return number if np.isfinite(number) else None


def description_with_uniscore(description, score):
    """Optionally put an available spectrum score below a panel's description."""
    if not SHOW_UNISCORE_LABELS:
        return description
    score = _finite_number(score)
    if score is None:
        return description
    return f"{description}\nUniScore: {score:.2f}"


def _reference_for_peak(reference_mass, peak_mass):
    if reference_mass is None:
        return None
    if np.isscalar(reference_mass):
        return _finite_number(reference_mass)
    references = np.asarray(list(reference_mass), dtype=float)
    references = references[np.isfinite(references)]
    if references.size == 0:
        return None
    return float(references[np.argmin(np.abs(references - peak_mass))])


def _config_values(engine, settings):
    values = {}
    config = getattr(engine, "config", None) if engine is not None else None
    for field in CONFIG_FIELDS:
        value = settings.get(field)
        if config is not None and hasattr(config, field):
            value = getattr(config, field)
        if isinstance(value, np.generic):
            value = value.item()
        elif isinstance(value, Path):
            value = str(value)
        values[f"setting_{field}"] = value
    for key, value in settings.items():
        column = key if key.startswith("setting_") else f"setting_{key}"
        if column not in values or values[column] is None:
            if isinstance(value, np.generic):
                value = value.item()
            elif isinstance(value, Path):
                value = str(value)
            values[column] = value
    return values


def panel_metrics(panel, label, engine=None, reference_mass=None, settings=None,
                  panel_type="deconvolved mass spectrum", notes=""):
    """Build one workbook row for a labeled figure panel."""
    settings = {} if settings is None else dict(settings)
    row = {
        "panel": panel,
        "label": label,
        "panel_type": panel_type,
        "reference_mass_Da": None,
        "reference_basis": settings.pop("reference_basis", ""),
        "apex_mass_Da": None,
        "apex_error_Da": None,
        "apex_abs_error_Da": None,
        "apex_error_ppm": None,
        "centroid_mass_Da": None,
        "centroid_error_Da": None,
        "centroid_abs_error_Da": None,
        "centroid_error_ppm": None,
        "DScore": None,
        "main_peak_area": None,
        "main_peak_area_window_low_Da": None,
        "main_peak_area_window_high_Da": None,
        "spectrum_area": None,
        "main_peak_area_fraction": None,
        "artifact_area_fraction": None,
        "main_peak_height_fraction": None,
        "FWHM_Da": None,
        "mass_resolving_power": None,
        "detected_peak_count": None,
        "UniScore": None,
        "fit_R_squared": None,
        "status": "metrics unavailable for this panel type",
        "notes": notes,
    }
    row.update(_config_values(engine, settings))
    if engine is None:
        if reference_mass is not None and np.isscalar(reference_mass):
            row["reference_mass_Da"] = _finite_number(reference_mass)
        return row

    peaks = list(getattr(getattr(engine, "pks", None), "peaks", []))
    if not peaks:
        row["status"] = "no peaks detected"
        row["detected_peak_count"] = 0
        return row

    main_peak = max(peaks, key=lambda peak: peak.height)
    apex = _finite_number(main_peak.mass)
    centroid = _finite_number(main_peak.centroid)
    if centroid == 0:
        centroid = None
    reference = _reference_for_peak(reference_mass, apex)
    dscore = _finite_number(main_peak.dscore)
    fwhm = _finite_number(main_peak.errorFWHM)
    mass_data = np.asarray(engine.data.massdat, dtype=float)
    spectrum_area = None
    main_area = None
    area_window = None
    if mass_data.ndim == 2 and mass_data.shape[0] > 1 and mass_data.shape[1] >= 2:
        spectrum_area = _finite_number(np.trapezoid(mass_data[:, 1], x=mass_data[:, 0]))
        if apex is not None and fwhm is not None and fwhm > 0:
            area_window = (apex - fwhm, apex + fwhm)
        else:
            integral_range = getattr(main_peak, "integralrange", [])
            if len(integral_range) == 2 and integral_range[1] > integral_range[0]:
                area_window = tuple(integral_range)
        if area_window is not None:
            in_window = ((mass_data[:, 0] >= area_window[0])
                         & (mass_data[:, 0] <= area_window[1]))
            window_data = mass_data[in_window]
            if len(window_data) > 1:
                main_area = _finite_number(
                    np.trapezoid(window_data[:, 1], x=window_data[:, 0])
                )
    if main_area is None:
        main_area = _finite_number(main_peak.integral)

    heights = np.asarray([peak.height for peak in peaks], dtype=float)
    height_sum = np.sum(heights)
    area_fraction = None
    if main_area is not None and spectrum_area is not None and spectrum_area > 0:
        area_fraction = main_area / spectrum_area

    row.update({
        "reference_mass_Da": reference,
        "apex_mass_Da": apex,
        "centroid_mass_Da": centroid,
        "DScore": dscore,
        "main_peak_area": main_area,
        "main_peak_area_window_low_Da": None if area_window is None else area_window[0],
        "main_peak_area_window_high_Da": None if area_window is None else area_window[1],
        "spectrum_area": spectrum_area,
        "main_peak_area_fraction": area_fraction,
        "artifact_area_fraction": None if area_fraction is None else 1 - area_fraction,
        "main_peak_height_fraction": None if height_sum <= 0 else float(main_peak.height / height_sum),
        "FWHM_Da": fwhm,
        "mass_resolving_power": None if not fwhm or apex is None else apex / fwhm,
        "detected_peak_count": len(peaks),
        "UniScore": _finite_number(getattr(engine.pks, "uniscore", None)),
        "fit_R_squared": _finite_number(getattr(engine.config, "error", None)),
        "status": "ok" if reference is not None else "reference mass not defined; accuracy unavailable",
    })
    if reference is not None and reference != 0:
        if apex is not None:
            error = apex - reference
            row.update({
                "apex_error_Da": error,
                "apex_abs_error_Da": abs(error),
                "apex_error_ppm": error / reference * 1e6,
            })
        if centroid is not None:
            error = centroid - reference
            row.update({
                "centroid_error_Da": error,
                "centroid_abs_error_Da": abs(error),
                "centroid_error_ppm": error / reference * 1e6,
            })
    return row


def write_metrics_workbook(output_path, rows):
    """Write panel metrics and their definitions to a formatted Excel workbook."""
    rows = list(rows)
    if not rows:
        raise ValueError("At least one panel metrics row is required")

    fixed_columns = list(rows[0])
    extra_columns = sorted({key for row in rows for key in row} - set(fixed_columns))
    columns = fixed_columns + extra_columns

    workbook = Workbook()
    sheet = workbook.active
    sheet.title = "Metrics"
    sheet.append(columns)
    for row in rows:
        sheet.append([row.get(column) for column in columns])
    sheet.freeze_panes = "A2"
    sheet.auto_filter.ref = sheet.dimensions
    for cell in sheet[1]:
        cell.font = Font(bold=True)
    for index, column in enumerate(columns, start=1):
        values = [str(column)] + ["" if row.get(column) is None else str(row.get(column)) for row in rows]
        sheet.column_dimensions[get_column_letter(index)].width = min(max(map(len, values)) + 2, 42)
        if column in COLOR_SCALED_METRICS:
            column_letter = get_column_letter(index)
            sheet.conditional_formatting.add(
                f"{column_letter}2:{column_letter}{len(rows) + 1}",
                ColorScaleRule(
                    start_type="min",
                    start_color="F8696B",
                    mid_type="percent",
                    mid_value=50,
                    mid_color="FFEB84",
                    end_type="max",
                    end_color="63BE7B",
                ),
            )

    definitions = workbook.create_sheet("Definitions")
    definitions.append(["field", "definition"])
    for field, definition in METRIC_DEFINITIONS:
        definitions.append([field, definition])
    definitions.append(["accuracy convention", "Signed errors are measured value minus reference value."])
    definitions.append(["blank cells", "The metric is not applicable or could not be calculated; see status and notes."])
    for cell in definitions[1]:
        cell.font = Font(bold=True)
    definitions.column_dimensions["A"].width = 30
    definitions.column_dimensions["B"].width = 110

    output_path = Path(output_path)
    output_path.parent.mkdir(parents=True, exist_ok=True)
    workbook.save(output_path)
    return output_path
