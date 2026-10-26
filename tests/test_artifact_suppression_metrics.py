import tempfile
import unittest
from pathlib import Path
from types import SimpleNamespace

import numpy as np
from openpyxl import load_workbook
from openpyxl.utils import get_column_letter

from PublicScripts.ArtifactSuppression import metrics
from PublicScripts.ArtifactSuppression.metrics import panel_metrics, write_metrics_workbook


class TestArtifactSuppressionMetrics(unittest.TestCase):
    def test_uniscore_plot_labels_are_optional_and_off_by_default(self):
        self.assertEqual(metrics.description_with_uniscore("Top 1", 0.75), "Top 1")

        metrics.SHOW_UNISCORE_LABELS = True
        try:
            self.assertEqual(
                metrics.description_with_uniscore("Top 1", 0.75),
                "Top 1\nUniScore: 0.75",
            )
        finally:
            metrics.SHOW_UNISCORE_LABELS = False

    def test_main_peak_metrics_and_nearest_reference(self):
        minor = SimpleNamespace(
            mass=100.0, centroid=101.0, height=2.0, integral=4.0,
            dscore=0.4, errorFWHM=4.0,
        )
        main = SimpleNamespace(
            mass=201.0, centroid=202.0, height=8.0, integral=12.0,
            dscore=0.9, errorFWHM=5.0,
        )
        engine = SimpleNamespace(
            config=SimpleNamespace(beta=10, suppression_topn=1),
            data=SimpleNamespace(massdat=np.array([
                [190.0, 0.0], [196.0, 1.0], [201.0, 2.0], [206.0, 1.0], [212.0, 0.0],
            ])),
            pks=SimpleNamespace(peaks=[minor, main], uniscore=0.75),
        )

        row = panel_metrics("B", "Top 1", engine, reference_mass=[100.0, 200.0])

        self.assertEqual(row["reference_mass_Da"], 200.0)
        self.assertEqual(row["apex_error_Da"], 1.0)
        self.assertEqual(row["centroid_error_Da"], 2.0)
        self.assertEqual(row["DScore"], 0.9)
        self.assertAlmostEqual(row["main_peak_area_fraction"], 15.0 / 21.0)
        self.assertAlmostEqual(row["artifact_area_fraction"], 6.0 / 21.0)
        self.assertEqual(row["main_peak_height_fraction"], 0.8)
        self.assertEqual(row["mass_resolving_power"], 40.2)
        self.assertEqual(row["setting_beta"], 10)
        self.assertEqual(row["setting_suppression_topn"], 1)

    def test_non_deconvolution_panel_has_explicit_status(self):
        row = panel_metrics("A", "Data", panel_type="input spectrum")
        self.assertIsNone(row["apex_mass_Da"])
        self.assertIn("unavailable", row["status"])

    def test_workbook_contains_every_row_and_definitions(self):
        rows = [
            panel_metrics("A", "Data", panel_type="input spectrum"),
            panel_metrics("B", "Grid", panel_type="2D deconvolution"),
        ]
        with tempfile.TemporaryDirectory() as directory:
            path = write_metrics_workbook(Path(directory, "Figure1.xlsx"), rows)
            workbook = load_workbook(path, data_only=True)

        self.assertEqual(workbook.sheetnames, ["Metrics", "Definitions"])
        self.assertEqual(workbook["Metrics"].max_row, 3)
        self.assertEqual(workbook["Metrics"]["A2"].value, "A")
        self.assertEqual(workbook["Metrics"]["A3"].value, "B")

    def test_color_scale_applies_only_to_calculated_non_mass_metrics(self):
        rows = [
            {"panel": "A", "apex_mass_Da": 100.0, "DScore": 0.2, "setting_beta": 0},
            {"panel": "B", "apex_mass_Da": 200.0, "DScore": 0.8, "setting_beta": 10},
        ]
        with tempfile.TemporaryDirectory() as directory:
            path = write_metrics_workbook(Path(directory, "colors.xlsx"), rows)
            sheet = load_workbook(path)["Metrics"]

        headers = [cell.value for cell in sheet[1]]
        dscore_column = get_column_letter(headers.index("DScore") + 1)
        apex_column = get_column_letter(headers.index("apex_mass_Da") + 1)
        beta_column = get_column_letter(headers.index("setting_beta") + 1)
        formatted_ranges = {str(item.sqref) for item in sheet.conditional_formatting}
        self.assertIn(f"{dscore_column}2:{dscore_column}3", formatted_ranges)
        self.assertNotIn(f"{apex_column}2:{apex_column}3", formatted_ranges)
        self.assertNotIn(f"{beta_column}2:{beta_column}3", formatted_ranges)

        rules = sheet.conditional_formatting[f"{dscore_column}2:{dscore_column}3"]
        self.assertEqual(len(rules), 1)
        self.assertEqual(rules[0].type, "colorScale")
        self.assertEqual([value.type for value in rules[0].colorScale.cfvo],
                         ["min", "percent", "max"])
        self.assertEqual([color.rgb for color in rules[0].colorScale.color],
                         ["00F8696B", "00FFEB84", "0063BE7B"])


if __name__ == "__main__":
    unittest.main()
